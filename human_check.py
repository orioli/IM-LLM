#!/usr/bin/env python3
"""
human_check.py: a small blind human check of the LLM judge, focused on the top tail.

  python human_check.py sample                 # picks ~100 ideas, writes the blind rating sheet + hidden key
  python human_check.py agree human_check/rater_*.xlsx   # after rating: agreement between raters and judge

Why not just the top ideas? To show the judge can tell great ideas from ordinary ones, raters must see
both. The sample is stratified by the judge's score: every band is covered, the top bands are
over-sampled (that is where the paper's claims are), and within each band the picks rotate across
conditions and cap ideas per need, so one theme (e.g. solid-state RF for M13) cannot fill a band.

Raters see only the need, the idea's title, its new elements and its rationale. Condition, effort,
judge scores and node IDs are removed; the key that links items back is a separate file.
"""
from __future__ import annotations

import argparse
import csv
import random
import re
from collections import defaultdict
from pathlib import Path

RESULTS = Path("results")
OUT = Path("human_check")
SEED = 11
OPERATOR = "SUBSTITUTE (replace parts with something cheaper or simpler)"

# (band name, lowest judge score, highest (exclusive), ideas to draw). Scores move in 0.25 steps.
BANDS = [("top", 4.25, 9.0, 30),      # top 2.4% of all ideas: 68 exist, so about half are rated
         ("high", 4.0, 4.25, 30),     # the 4.0 band; with "top" this is the paper's top-idea threshold
         ("mid", 3.5, 4.0, 25),
         ("low", 0.0, 3.5, 15)]
MAX_PER_NEED = 3                      # per band: at most this many ideas for the same need
CONDITIONS = ["A", "B", "R", "C"]
CRITERIA = ["relevance", "novelty", "feasibility", "specificity"]
CRITERIA_TEXT = {"relevance": "how well it serves the target need",
                 "novelty": "how radical and non-obvious it is",
                 "feasibility": "how technically plausible it is",
                 "specificity": "how concrete it is (named components, mechanisms)"}
PURSUE = ["yes", "maybe", "no"]
NODE_ID = re.compile(r"\(?\b[ACDFMPTX]\d{1,2}\b\)?")


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def clean(text: str) -> str:
    """Remove graph node IDs (they could hint at the condition) and tidy spacing."""
    text = NODE_ID.sub("", text or "")
    return re.sub(r"\s{2,}", " ", re.sub(r"\s+([,.;:])", r"\1", text)).strip()


def elements(text: str) -> str:
    """'Component: X — d; Technology: Y — d' -> one line per element."""
    parts = [clean(p) for p in (text or "").split("; ") if p.strip()]
    return "\n".join(f"• {p}" for p in parts)


# ── sample ──────────────────────────────────────────────────────────
def cmd_sample(args) -> None:
    rng = random.Random(args.seed)
    ideas = [r for r in read_csv(args.ideas) if num(r.get("score")) is not None]
    picked = []
    for band, lo, hi, n in BANDS:
        pool = defaultdict(list)                       # condition -> ideas in this band
        for r in ideas:
            if lo <= num(r["score"]) < hi:
                pool[r["condition"]].append(r)
        for c in pool:
            rng.shuffle(pool[c])
        per_need, got = defaultdict(int), []
        order = [c for c in CONDITIONS if pool.get(c)]
        # round-robin over conditions, so scarce conditions (R, C in the top band) are all included
        while len(got) < n and any(pool[c] for c in order):
            for c in order:
                while pool[c]:
                    r = pool[c].pop()
                    if per_need[r["target_need"]] < MAX_PER_NEED:
                        per_need[r["target_need"]] += 1
                        got.append({**r, "band": band})
                        break
                if len(got) >= n:
                    break
        picked += got
        counts = {c: sum(g["condition"] == c for g in got) for c in CONDITIONS}
        print(f"{band:5} {lo:g}–{hi if hi < 9 else 5:g}: {len(got):3} ideas  "
              + "  ".join(f"{c}={k}" for c, k in counts.items()))
    rng.shuffle(picked)
    for i, r in enumerate(picked, 1):
        r["item"] = f"H{i:03d}"

    args.out.mkdir(parents=True, exist_ok=True)
    key_cols = ["item", "band", "idea_id", "run_id", "condition", "effort", "target_need", "score"] + CRITERIA
    with (args.out / "key.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=key_cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(picked)

    sheet = [{"item": r["item"], "need": r["need_label"], "idea": clean(r["title"]),
              "what it adds": elements(r["new_elements"]), "why it serves the need": clean(r["rationale"])}
             for r in picked]
    write_sheet(args.out / "rating_sheet.xlsx", sheet)
    print(f"{len(picked)} ideas -> {args.out / 'rating_sheet.xlsx'} (give to raters), "
          f"{args.out / 'key.csv'} (keep hidden)")


def write_sheet(path: Path, rows: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    lines = ["Rating microwave innovation ideas", "",
             "Product: a household microwave oven.",
             f"Each idea was generated for one market need, using the innovation operator {OPERATOR}.",
             "The brief asked for radical, wild ideas that serve the target need.", "",
             "Rate EACH idea on its own, from 1 (poor) to 5 (excellent):"]
    lines += [f"  • {c}: {t}" for c, t in CRITERIA_TEXT.items()]
    lines += ["", "Pursue: would you put this idea forward for further development? yes / maybe / no",
              "Comment is optional.", "",
              "Ignore writing style and length. Ideas are in random order; please rate all of them,",
              "and work alone — do not discuss ratings with other raters."]
    for i, line in enumerate(lines, 1):
        ins.cell(row=i, column=1, value=line)
    ins["A1"].font = Font(bold=True, size=14)
    ins.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Ideas")
    cols = ["item", "need", "idea", "what it adds", "why it serves the need"] + CRITERIA + ["pursue", "comment"]
    widths = [7, 26, 30, 60, 45] + [11] * len(CRITERIA) + [9, 30]
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    head = PatternFill("solid", fgColor="E1E0D9")
    for j, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=j)
        cell.font, cell.fill = Font(bold=True), head
        ws.column_dimensions[cell.column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    n = len(rows) + 1
    score = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                           showErrorMessage=True, error="Enter a whole number from 1 to 5")
    pursue = DataValidation(type="list", formula1='"' + ",".join(PURSUE) + '"', allow_blank=True)
    ws.add_data_validation(score)
    ws.add_data_validation(pursue)
    first = cols.index(CRITERIA[0]) + 1
    for j in range(first, first + len(CRITERIA)):
        letter = ws.cell(row=1, column=j).column_letter
        score.add(f"{letter}2:{letter}{n}")
    letter = ws.cell(row=1, column=cols.index("pursue") + 1).column_letter
    pursue.add(f"{letter}2:{letter}{n}")
    wb.save(path)


# ── agree ───────────────────────────────────────────────────────────
def read_ratings(path: Path) -> dict:
    """item -> {criterion: score, 'pursue': ...} from a filled sheet (.xlsx or .csv)."""
    if path.suffix == ".xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True)["Ideas"]
        rows = list(ws.iter_rows(values_only=True))
        head = [str(h) for h in rows[0]]
        recs = [dict(zip(head, r)) for r in rows[1:]]
    else:
        recs = read_csv(path)
    out = {}
    for r in recs:
        vals = {c: num(r.get(c)) for c in CRITERIA}
        if r.get("item") and all(v is not None for v in vals.values()):
            out[str(r["item"])] = {**vals, "pursue": str(r.get("pursue") or "").strip().lower()}
    return out


def ranks(xs):
    order = sorted(range(len(xs)), key=lambda i: xs[i])
    r = [0.0] * len(xs)
    i = 0
    while i < len(xs):
        j = i
        while j + 1 < len(xs) and xs[order[j + 1]] == xs[order[i]]:
            j += 1
        for k in range(i, j + 1):
            r[order[k]] = (i + j) / 2 + 1
        i = j + 1
    return r


def pearson(x, y):
    mx, my = sum(x) / len(x), sum(y) / len(y)
    sxy = sum((a - mx) * (b - my) for a, b in zip(x, y))
    sx = sum((a - mx) ** 2 for a in x) ** 0.5
    sy = sum((b - my) ** 2 for b in y) ** 0.5
    return sxy / (sx * sy) if sx and sy else float("nan")


def spearman(x, y):
    return pearson(ranks(x), ranks(y))


def kripp_alpha_interval(units: list[list[float]]) -> float:
    """Krippendorff's alpha, interval metric; units = ratings per item (items with < 2 ratings ignored)."""
    units = [u for u in units if len(u) >= 2]
    n = sum(len(u) for u in units)
    if n < 2:
        return float("nan")
    d_o = sum(sum((a - b) ** 2 for a in u for b in u) / (len(u) - 1) for u in units) / n
    allv = [v for u in units for v in u]
    d_e = sum((a - b) ** 2 for a in allv for b in allv) / (n * (n - 1))
    return 1 - d_o / d_e if d_e else float("nan")


def auc(pos, neg):
    """Chance a random 'pos' idea gets a higher human score than a random 'neg' one (ties count half)."""
    if not pos or not neg:
        return float("nan")
    return sum((p > q) + 0.5 * (p == q) for p in pos for q in neg) / (len(pos) * len(neg))


def cmd_agree(args) -> None:
    key = {r["item"]: r for r in read_csv(args.key)}
    raters = {p.stem: read_ratings(p) for p in args.ratings}
    for name, rs in raters.items():
        print(f"{name}: {len(rs)} of {len(key)} items rated")
    items = [i for i in key if any(i in rs for rs in raters.values())]
    hum = {i: [sum(rs[i][c] for c in CRITERIA) / len(CRITERIA) for rs in raters.values() if i in rs] for i in items}
    hmean = {i: sum(v) / len(v) for i, v in hum.items()}

    lines = ["# Human check of the LLM judge", "",
             f"{len(items)} ideas, {len(raters)} rater(s). Human score = mean of the four criteria (1–5); "
             "judge score = the blind individual score from `top_ideas.py`.", ""]
    if len(raters) >= 2:
        a = kripp_alpha_interval([hum[i] for i in items])
        names = list(raters)
        pairs = []
        for x in range(len(names)):
            for y in range(x + 1, len(names)):
                common = [i for i in items if i in raters[names[x]] and i in raters[names[y]]]
                sx = [sum(raters[names[x]][i][c] for c in CRITERIA) for i in common]
                sy = [sum(raters[names[y]][i][c] for c in CRITERIA) for i in common]
                pairs.append(f"{names[x]}–{names[y]} ρ = {spearman(sx, sy):.2f}")
        lines += ["## Agreement between raters", "",
                  f"- Krippendorff's α (interval) on the idea score: **{a:.2f}** "
                  "(≥ 0.67 acceptable, ≥ 0.80 good)",
                  "- Spearman between raters: " + "; ".join(pairs), ""]

    js = [num(key[i]["score"]) for i in items]
    hs = [hmean[i] for i in items]
    lines += ["## Judge vs humans", "",
              f"- Spearman, judge score vs human mean score: **ρ = {spearman(js, hs):.2f}** (n = {len(items)})"]
    for c in CRITERIA:
        jc = [num(key[i][c]) for i in items]
        hc = [sum(rs[i][c] for rs in raters.values() if i in rs) / len(hum[i]) for i in items]
        ok = [(a, b) for a, b in zip(jc, hc) if a is not None]
        lines.append(f"  - {c}: ρ = {spearman([a for a, _ in ok], [b for _, b in ok]):.2f}")
    top = [hmean[i] for i in items if num(key[i]["score"]) >= 4.0]
    rest = [hmean[i] for i in items if num(key[i]["score"]) < 4.0]
    lines += ["", f"- **Top-idea check:** an idea the judge scores ≥ 4.0 gets a higher human score than one it "
              f"scores below 4.0 with probability **{auc(top, rest):.2f}** (0.5 = chance, 1 = always)", ""]

    lines += ["## Human score by judge band", "",
              "| Judge band | Ideas | Human score (mean) | Pursue = yes |", "| --- | --- | --- | --- |"]
    for band, lo, hi, _ in BANDS:
        sel = [i for i in items if key[i]["band"] == band]
        if not sel:
            continue
        yes = [rs[i]["pursue"] == "yes" for rs in raters.values() for i in sel if i in rs and rs[i]["pursue"]]
        lines.append(f"| {band} ({lo:g}–{hi if hi < 9 else 5:g}) | {len(sel)} | "
                     f"{sum(hmean[i] for i in sel) / len(sel):.2f} | "
                     + (f"{sum(yes) / len(yes):.0%}" if yes else "—") + " |")
    lines += ["", "Human scores should fall from band to band if the judge ranks ideas the way people do.", "",
              "## Human score by condition (descriptive; the sample is not balanced for this)", "",
              "| Condition | Ideas | Human score (mean) | Judge score (mean) |", "| --- | --- | --- | --- |"]
    for c in CONDITIONS:
        sel = [i for i in items if key[i]["condition"] == c]
        if sel:
            lines.append(f"| {c} | {len(sel)} | {sum(hmean[i] for i in sel) / len(sel):.2f} | "
                         f"{sum(num(key[i]['score']) for i in sel) / len(sel):.2f} |")

    out = args.key.parent / "agreement.md"
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    with (args.key.parent / "human_scores.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["item", "band", "condition", "effort", "target_need", "judge_score", "human_score", "n_raters"])
        for i in items:
            k = key[i]
            w.writerow([i, k["band"], k["condition"], k["effort"], k["target_need"], k["score"],
                        round(hmean[i], 3), len(hum[i])])
    print("\n".join(lines))
    print(f"\n-> {out}, {args.key.parent / 'human_scores.csv'}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sample", help="draw the stratified blind sample and write the rating sheet")
    s.add_argument("--ideas", type=Path, default=RESULTS / "ideas_ranked.csv")
    s.add_argument("--out", type=Path, default=OUT)
    s.add_argument("--seed", type=int, default=SEED)
    s.set_defaults(func=cmd_sample)
    a = sub.add_parser("agree", help="agreement between raters and the judge")
    a.add_argument("ratings", nargs="+", type=Path, help="filled rating sheets, one per rater (.xlsx or .csv)")
    a.add_argument("--key", type=Path, default=OUT / "key.csv")
    a.set_defaults(func=cmd_agree)
    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
