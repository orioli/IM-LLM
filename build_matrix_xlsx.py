import yaml, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
g = yaml.safe_load(open(sys.argv[1])); out = sys.argv[2]
nodes = {n['id']: n for n in g['nodes']}; edges = g['edges']
num = lambda i: int(''.join(c for c in i if c.isdigit()) or 0)
ORDER = ['Field','Technology','Component','Function','MarketNeed','Actor','Constraint','Product']
DCOL = dict(Product='4C72B0', Component='55A868', Function='C44E52', MarketNeed='8172B2', Actor='CCB974',
            Constraint='64B5CD', Technology='DD8452', Field='937860')
LIGHT = dict(Product='DCE4F2', Component='DDEFE1', Function='F5DCDD', MarketNeed='E6E2F1', Actor='F3EEDA',
             Constraint='DCEFF5', Technology='F8E3D6', Field='E9E2DB')
SYM = {  # edge type -> (symbol, cell fill)
 'has_component':('●','DDE3EA'), 'enabled_by':('●','F8E3D6'), 'rooted_in':('●','E9E2DB'),
 'performs_function':('●','DDEFE1'), 'addresses_need':('●','E6E2F1'), 'enables_function':('→','F5DCDD'),
 'feeds':('→','DDEFE1'), 'actuates':('⇒','DDEFE1'), 'relevant_to':('·','F1F3F5'),
 'causes':('▲','FFD8D8'), 'contradicts':('✕','FF9E9E'), 'partially_mitigates':('◐','C9EFD0'), 'resolves':('★','FFE8A3')}
F = 'Arial'
thin = Side(style='thin', color='D0D5DA'); thick = Side(style='medium', color='495057')
def ids(t): return sorted([i for i,n in nodes.items() if n['type']==t], key=num)
ALL = [i for t in ORDER for i in ids(t)]
E = {}
for e in edges: E[(e['from'],e['to'])] = e['type']
wb = Workbook()

# ── README
ws = wb.active; ws.title = 'README'
rows = [
 ('Innovation matrix — basic countertop microwave (graph G as matrices)', True),
 (f"Source: {g['meta']['graph_id']} · {len(nodes)} nodes · {len(edges)} edges", False),
 ('', False),
 ('Sheets', True),
 ('MDM (full)  — Multi-Domain Matrix: every node × every node. Row = FROM node, column = TO node. Coloured blocks = domain pairs.', False),
 ('Comp×Function — functional analysis matrix (which component performs which function).', False),
 ('Function×Need — QFD-style: which function addresses which market need.', False),
 ('Comp×Tech, Tech×Field — technology roots of each component (lower layers of the Bernegger matrix).', False),
 ('Comp DSM — component × component energy/mechanical flow (feeds / actuates).', False),
 ('Constraints — every constraint against every node it touches, with the relation type.', False),
 ('Actor×Need — who cares about what.', False),
 ('Nodes, Edges — the raw lists (one row each), handy for filters and pivots.', False),
 ('', False),
 ('Cell symbols', True)]
for t,(s,_) in SYM.items(): rows.append((f'{s}   {t}', False))
rows += [('', False), ('Domain-mapping sheets ignore edge direction: a mark means the two nodes are linked. Totals rows/columns are COUNTA formulas.', False)]
for r,(txt,b) in enumerate(rows,1):
    c = ws.cell(r,1,txt); c.font = Font(name=F, bold=b, size=14 if r==1 else 10)
for r,(t,(s,fill)) in enumerate(SYM.items(), 15):
    ws.cell(r,1).fill = PatternFill('solid', fgColor=fill)
ws.column_dimensions['A'].width = 120

def matrix(title, R, C, directed, note):
    ws = wb.create_sheet(title)
    ws.cell(1,1,note).font = Font(name=F, italic=True, size=9, color='666666')
    HR, HC = 3, 3   # header row / first data col
    ws.cell(HR,1,'ID').font = Font(name=F,bold=True,size=9)
    ws.cell(HR,2,'Label').font = Font(name=F,bold=True,size=9)
    for j,cid in enumerate(C):
        col = HC+j; t = nodes[cid]['type']
        c = ws.cell(HR-1,col,cid); c.font = Font(name=F,bold=True,size=8,color='FFFFFF')
        c.fill = PatternFill('solid',fgColor=DCOL[t]); c.alignment = Alignment(horizontal='center')
        c = ws.cell(HR,col,nodes[cid]['label']); c.font = Font(name=F,size=8)
        c.alignment = Alignment(text_rotation=90, vertical='bottom', horizontal='center', wrap_text=True)
        c.fill = PatternFill('solid',fgColor=LIGHT[t])
        ws.column_dimensions[CL(col)].width = 3.6
    ws.row_dimensions[HR].height = 190
    last_c = HC+len(C)-1; tot_c = last_c+1
    ws.cell(HR,tot_c,'Total').font = Font(name=F,bold=True,size=9)
    ws.cell(HR,tot_c).alignment = Alignment(text_rotation=90, vertical='bottom', horizontal='center')
    ws.column_dimensions[CL(tot_c)].width = 5
    for i,rid in enumerate(R):
        r = HR+1+i; t = nodes[rid]['type']
        c = ws.cell(r,1,rid); c.font = Font(name=F,bold=True,size=8,color='FFFFFF'); c.fill = PatternFill('solid',fgColor=DCOL[t])
        c = ws.cell(r,2,nodes[rid]['label']); c.font = Font(name=F,size=8); c.fill = PatternFill('solid',fgColor=LIGHT[t])
        for j,cid in enumerate(C):
            col = HC+j; cell = ws.cell(r,col)
            et = E.get((rid,cid)) or (None if directed else E.get((cid,rid)))
            # block shading for MDM
            if directed and nodes[rid]['type']==nodes[cid]['type']:
                cell.fill = PatternFill('solid',fgColor='F7F7F7')
            if rid == cid: cell.fill = PatternFill('solid',fgColor='495057')
            if et:
                s, fill = SYM.get(et, ('?','EEEEEE'))
                cell.value = s; cell.fill = PatternFill('solid',fgColor=fill)
                cell.font = Font(name=F,size=9,bold=et in('contradicts','causes'),color='B22222' if et in('contradicts','causes') else '222222')
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = Border(left=thin,right=thin,top=thin,bottom=thin)
        ws.cell(r,tot_c,f'=COUNTA({CL(HC)}{r}:{CL(last_c)}{r})').font = Font(name=F,size=8,bold=True)
        ws.cell(r,tot_c).alignment = Alignment(horizontal='center')
    last_r = HR+len(R); tr = last_r+1
    ws.cell(tr,2,'Total').font = Font(name=F,bold=True,size=9)
    for j in range(len(C)+1):
        col = HC+j
        ws.cell(tr,col,f'=COUNTA({CL(col)}{HR+1}:{CL(col)}{last_r})' if j<len(C) else f'=SUM({CL(HC)}{tr}:{CL(last_c)}{tr})').font = Font(name=F,size=8,bold=True)
        ws.cell(tr,col).alignment = Alignment(horizontal='center')
    # thick borders between domain blocks
    def bounds(L):
        b=[]; 
        for k in range(1,len(L)):
            if nodes[L[k]]['type']!=nodes[L[k-1]]['type']: b.append(k)
        return b
    for k in bounds(C):
        col = HC+k
        for r in range(HR-1, tr):
            c = ws.cell(r,col); c.border = Border(left=thick,right=c.border.right,top=c.border.top,bottom=c.border.bottom)
    for k in bounds(R):
        r = HR+1+k
        for col in range(1, tot_c):
            c = ws.cell(r,col); c.border = Border(top=thick,left=c.border.left,right=c.border.right,bottom=c.border.bottom)
    ws.column_dimensions['A'].width = 5; ws.column_dimensions['B'].width = 46
    ws.freeze_panes = ws.cell(HR+1, HC)
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1; ws.sheet_properties.pageSetUpPr.fitToPage=True
    return ws

matrix('MDM (full)', ALL, ALL, True, 'Multi-Domain Matrix. Row = FROM, column = TO. Grey cells = same-domain block; dark diagonal = self. Thick lines separate domains.')
matrix('Comp×Function', ids('Component'), ids('Function'), False, 'Functional analysis matrix: ● = component performs function.')
matrix('Function×Need', ids('Function'), ids('MarketNeed'), False, 'QFD-style matrix: ● = function addresses need.')
matrix('Comp×Tech', ids('Component'), ids('Technology'), False, '● = component is enabled by technology.')
matrix('Tech×Field', ids('Technology'), ids('Field'), False, '● = technology is rooted in field.')
matrix('Comp DSM', ids('Component'), ids('Component'), True, 'Design Structure Matrix. Row feeds/actuates column: → feeds (energy/signal), ⇒ actuates (mechanical), · relevant.')
others = [i for t in ['Technology','Component','Function','MarketNeed','Actor','Product'] for i in ids(t)]
matrix('Constraints', ids('Constraint'), others, False, '▲ causes · ✕ contradicts · ◐ partially mitigates · · relevant (direction ignored).')
matrix('Actor×Need', ids('Actor'), ids('MarketNeed')+ids('Constraint'), False, '· = actor is relevant to need / constraint.')

# ── Nodes / Edges lists
ws = wb.create_sheet('Nodes')
hdr = ['id','type','label','detail','proposed_by','status']
for j,h in enumerate(hdr,1): ws.cell(1,j,h).font = Font(name=F,bold=True)
for r,i in enumerate(ALL,2):
    n = nodes[i]
    for j,h in enumerate(hdr,1): ws.cell(r,j,n.get(h,'')).font = Font(name=F,size=9)
    ws.cell(r,1).fill = PatternFill('solid',fgColor=LIGHT[n['type']])
for col,w in zip('ABCDEF',[6,12,60,80,14,10]): ws.column_dimensions[col].width = w
ws.auto_filter.ref = f'A1:F{len(ALL)+1}'; ws.freeze_panes = 'A2'
ws = wb.create_sheet('Edges')
hdr = ['from','from_label','type','to','to_label']
for j,h in enumerate(hdr,1): ws.cell(1,j,h).font = Font(name=F,bold=True)
for r,e in enumerate(edges,2):
    vals = [e['from'], nodes[e['from']]['label'], e['type'], e['to'], nodes[e['to']]['label']]
    for j,v in enumerate(vals,1): ws.cell(r,j,v).font = Font(name=F,size=9)
for col,w in zip('ABCDE',[6,50,18,6,50]): ws.column_dimensions[col].width = w
ws.auto_filter.ref = f'A1:E{len(edges)+1}'; ws.freeze_panes = 'A2'
wb.save(out); print('saved')
